# -*- coding: utf-8 -*-
"""将 thesis/毕业论文模拟稿.md 转换为带封面与目录的 Word 文档"""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT

MD = "thesis/毕业论文模拟稿.md"
OUT = "thesis/毕业论文模拟稿.docx"

doc = Document()
style = doc.styles["Normal"]
style.font.name = "宋体"
style.font.size = Pt(12)
style.element.rPr.rFonts.set(__import__("docx").oxml.ns.qn("w:eastAsia"), "宋体")

def add_heading(text, level):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.name = "黑体"
    run._element.rPr.rFonts.set(__import__("docx").oxml.ns.qn("w:eastAsia"), "黑体")
    run.font.size = Pt(16 if level == 1 else (14 if level == 2 else 12))
    run.font.bold = True
    return h

def add_table(rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(set(c) <= set("-: ") for c in r)]
    t = doc.add_table(rows=len(cells), cols=len(cells[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(cells):
        for j, val in enumerate(row):
            p = t.cell(i, j).paragraphs[0]
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            run.font.bold = (i == 0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ---------- 封面 ----------
for _ in range(6):
    doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("毕业论文（预演）模拟稿"); r.font.size = Pt(22); r.font.bold = True
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("管理层语调的NLP测度构建及其对企业未来业绩的预测力研究\n——基于A股上市公司年报MD&A文本的五步流水线实证")
r.font.size = Pt(16); r.font.bold = True
for _ in range(4):
    doc.add_paragraph()
for line in ["作者：马傲君", "专业：大数据管理与应用", "完成日期：2026年10月"]:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(line).font.size = Pt(14)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---------- 目录页 ----------
add_heading("目 录", 1)
toc_items = [
    "摘要", "1. 绪论", "　1.1 选题背景　1.2 研究意义　1.3 文献综述与研究缺口（Gap）",
    "2. 方法", "　2.1 数据来源与样本　2.2 文本预处理　2.3 测度构造　2.4 变量定义表　2.5 五步流水线技术路线",
    "3. 分析", "　3.1 描述性统计　3.2 信度检验　3.3 效度检验　3.4 回归分析　3.5 稳健性检验",
    "4. 结论", "　4.1 主要发现　4.2 局限性　4.3 未来展望",
    "参考文献", "附录：可复现代码与数据说明",
]
for it in toc_items:
    p = doc.add_paragraph(it)
    p.paragraph_format.space_after = Pt(4)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---------- 正文 ----------
lines = open(MD, encoding="utf-8").read().splitlines()
i = 0
in_toc = False
while i < len(lines):
    ln = lines[i]
    if ln.startswith("# ") and i == 0:
        i += 1; continue
    if ln.startswith("## 目录"):
        in_toc = True; i += 1; continue
    if in_toc:
        if ln.startswith("## "):
            in_toc = False
        else:
            i += 1; continue
    if ln.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            rows.append(lines[i]); i += 1
        add_table(rows); continue
    if ln.startswith("### "):
        add_heading(ln[4:].strip(), 3)
    elif ln.startswith("## "):
        add_heading(ln[3:].strip(), 2)
    elif ln.startswith("> "):
        p = doc.add_paragraph(ln[2:].strip())
        p.paragraph_format.left_indent = Cm(1)
        for r in p.runs: r.font.italic = False; r.font.size = Pt(11)
    elif ln.startswith("**") and ln.endswith("**") and len(ln) < 60:
        p = doc.add_paragraph(); r = p.add_run(ln.strip("*"))
        r.font.bold = True; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif ln.startswith("---"):
        pass
    elif ln.strip():
        p = doc.add_paragraph(ln.replace("**", ""))
        p.paragraph_format.first_line_indent = Pt(24)
        p.paragraph_format.line_spacing = 1.5
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    i += 1

doc.save(OUT)
print("已生成:", OUT)
